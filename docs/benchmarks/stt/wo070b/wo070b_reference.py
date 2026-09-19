"""WO-070B reference dataset builder.

Reads the COMPLETED WO-068 human-review artifact
(``WO-068-HUMAN-REVIEW-EXTENDED.xlsx``) READ-ONLY and emits the benchmark-only
reference CSV ``human_reference_transcripts.csv``.

Rules enforced here (WO-070B §5-§9):
  * the transcript comes ONLY from the sheet's human-entered ``transcript`` cell;
  * ``audible_voice = NO``  -> reference_status = NON_AUDIBLE, transcript = "";
  * audible + non-empty human transcript -> HUMAN_VERIFIED;
  * audible + empty human transcript    -> NO_HUMAN_TRANSCRIPT, transcript = "";
  * nothing is derived from STT output, callsign, notes, filename or inference.

The workbook is opened with openpyxl in read-only/data_only mode and is never
written back.
"""

import csv
import os

REFERENCE_COLUMNS = [
    "message_id",
    "stream_id",
    "reference_transcript",
    "reference_status",
    "reviewer",
    "review_time",
]

STATUS_HUMAN_VERIFIED = "HUMAN_VERIFIED"
STATUS_NON_AUDIBLE = "NON_AUDIBLE"
STATUS_NO_HUMAN_TRANSCRIPT = "NO_HUMAN_TRANSCRIPT"


def _cell_str(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat") and not isinstance(value, str):
        try:
            return str(value)
        except Exception:
            return ""
    return str(value).strip()


def build_reference_rows(xlsx_path):
    """Return the ordered list of reference row dicts read from the workbook."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        if "Human Review" not in wb.sheetnames:
            raise KeyError("sheet 'Human Review' missing from workbook")
        ws = wb["Human Review"]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    header = [(_cell_str(h)).strip() for h in rows[0]]
    if "message_id" not in header:
        raise KeyError("column 'message_id' missing from Human Review sheet")
    if "transcript" not in header:
        raise KeyError("column 'transcript' missing from Human Review sheet")

    out = []
    for raw in rows[1:]:
        if raw is None or all(c is None for c in raw):
            continue
        rec = {header[i]: raw[i] for i in range(len(header)) if i < len(raw)}
        message_id = _cell_str(rec.get("message_id"))
        if not message_id:
            continue

        audible = _cell_str(rec.get("audible_voice")).upper()
        transcript = _cell_str(rec.get("transcript"))
        reviewer = _cell_str(rec.get("reviewer"))
        review_time = _cell_str(rec.get("review_time"))

        if audible == "NO":
            status = STATUS_NON_AUDIBLE
            reference_transcript = ""
        elif transcript:
            status = STATUS_HUMAN_VERIFIED
            reference_transcript = transcript
        else:
            status = STATUS_NO_HUMAN_TRANSCRIPT
            reference_transcript = ""

        out.append(
            {
                "message_id": message_id,
                "stream_id": _cell_str(rec.get("stream_id")),
                "reference_transcript": reference_transcript,
                "reference_status": status,
                "reviewer": reviewer,
                "review_time": review_time,
            }
        )
    return out


def write_reference_csv(rows, out_path):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=REFERENCE_COLUMNS, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in REFERENCE_COLUMNS})
    return out_path


def read_reference_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))
